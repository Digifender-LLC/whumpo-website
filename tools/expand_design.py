from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import RGBColor
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'docs' / 'reference' / 'ai_managed_cloud_mvp_design.docx'
OUTPUT = Path(__file__).resolve().parents[1] / 'docs' / 'ai_managed_cloud_mvp_design_revised.docx'

doc = Document(SOURCE)
doc.paragraphs[0].style = doc.styles['Title']
doc.paragraphs[0].text = 'AI Managed Cloud Platform MVP Design'
for run in doc.paragraphs[0].runs:
    run.font.color.rgb = RGBColor(0, 0, 0)
for element in (doc.styles['Title']._element, doc.paragraphs[0]._p):
    ppr = element.find(qn('w:pPr'))
    if ppr is not None:
        border = ppr.find(qn('w:pBdr'))
        if border is not None:
            ppr.remove(border)
doc.paragraphs[1].text = 'Product and technical design'
doc.paragraphs[3].text = 'Version 0.5  |  September 29, 2026'

first_section = next(p for p in doc.paragraphs if p.text.startswith('1. Executive Summary'))
intro = first_section.insert_paragraph_before()
intro.add_run('Current implementation. ').bold = True
intro.add_run('Whumpo runs an invitation-only test at whumpo.com with passkeys, conversational onboarding, metered Muse edits, reusable websites, photos, publication and source export. Contact forms, directory CRUD and recovery, comparisons, health checks and reviewed integration controls are also deployed. Real customer provider setup and acceptance remain pending; the full MVP is incomplete. Sections 1 through 25 preserve the original product vision. Sections 26 through 33 supersede implementation and infrastructure assumptions with current evidence and remaining work. Stripe remains deferred and payments are disabled.')

def paragraph(text):
    doc.add_paragraph(text)

def heading(text):
    doc.add_heading(text, 1)

def sub(text):
    doc.add_heading(text, 2)

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(t.rows[0].cells, headers):
        cell.text = label
        for run in cell.paragraphs[0].runs: run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row): cell.text = value
    from docx.oxml import OxmlElement
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in t.rows:
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
    if len(headers) == 2:
        for row in t.rows:
            row.cells[0].width = Inches(1.9)
            row.cells[1].width = Inches(5.1)
    return t

heading('26 Deployed Test Platform and Scope')
paragraph('The homepage advertises the product. The signed-in workspace opens into a persistent conversation with compact project, live-site and usage statistics. A Details menu exposes files, versions, domains, connected accounts, security and billing. Chat supports guided commands and validated natural-language proposals. Provider consent and passkey device interaction occur outside the conversation; domain and backup reviews remain explicit customer decisions.')
paragraph('The infrastructure constraint is GitHub free features and Cloudflare services within existing included allowances. The owner authorized the current Workers Standard account allowance; this is not a separate Workers Free account. The owner subsequently authorized the $5/month Workers Paid resources; no further subscription upgrades are implied. The application is invitation-only, and whumpo.com is the platform domain; customer sites use whumpo.site subdomains. Whumpo is the product name.')
table(['Boundary', 'Deployed implementation'], [
 ('Workspace', 'whumpo-platform Worker, Static Assets and Turnstile at whumpo.com'),
 ('State', 'One D1 database for identity, chat, projects, usage, versions and publication reviews'),
 ('Files', 'Private Standard R2 bucket whumpo-files; public bucket access disabled'),
 ('Published sites', 'whumpo-sites Worker at PROJECT_ID.whumpo.site; current version and exact photo references only'),
 ('AI provider', 'OpenAI-compatible Responses transport to OpenCode Go Muse Spark 1.3 Contributor'),
 ('Source control', 'Private GitHub repository zemdregon/Whumpo; deployment from the local CLI'),
])
paragraph('Durable Objects, Workflows, Queues, Containers, Sandbox, Workers AI, AI Gateway, Artifacts and Registrar described earlier are architectural options, not deployed dependencies. The current assembler accepts reviewed components and bounded configuration; it does not execute arbitrary customer-generated code. General dynamic apps and isolated build execution remain open requirements under the cost constraint.')
paragraph('Source of truth: wrangler.jsonc, wrangler.sites.jsonc, migrations, src and the deployment log. The public Worker has D1 and R2 bindings, but no identity or AI secrets. These bindings are not read-only credentials; route code enforces the public read boundary. No customer receives platform provider credentials.')

heading('27 Identity Conversation and Action Authority')
sub('Identity and tenant ownership')
paragraph('Passkeys are the default registration and login method. SimpleWebAuthn validates the relying party, origin, challenge, signature, user verification and authenticator counter. Challenges are hashed, browser-cookie bound, single-use and five minutes long. Invitations and Turnstile gate signup. Optional passwords, rotating recovery codes, session revocation and recent authentication for security changes are available.')
paragraph('Every project resource is joined to the authenticated owner before it is returned or changed. Cross-owner resources return 404. Browser mutations check Origin; session cookies are Secure, HttpOnly and SameSite. Provider credentials are Worker secrets. opencode.key and local test credentials are ignored by Git and must never enter prompts, exports or logs.')
paragraph('Production passkeys are bound to whumpo.com. The migration to whumpo.com is complete. There were no registered production passkeys at cutover; no old-origin fallback or redirect is retained. Signed assertion fixtures pass live verification, but synced-device passkey prompts still require a user-device acceptance check.')
sub('Conversation consistency')
paragraph('Guided onboarding collects idea, project name, audience or visitor goal, and style preferences without inference, resumes after interruption, and creates the first private draft. Audience and style can be skipped. Private planning notes are not automatically exported or copied into page text. Each message carries a UUID and expected revision. Atomic D1 claims prevent competing turns and duplicate work. Pending model calls serialize the conversation. The test retains up to 2,000 turns and displays the latest 100.')
sub('The model proposes and the server decides')
paragraph('The model returns a validated JSON proposal in text rather than native SDK function calls. Server code validates arguments, ownership and quotas. Actions cover projects, private drafts, focused edits, owned inspection, static/data exports, domain proof checks, DNS discovery/reviews, backup/hosting reviews, health checks, contact activation reviews and directory CRUD/recovery reviews. The model cannot confirm reviews, grant entitlements, choose arbitrary provider destinations or execute customer backend code. Contact inbox contents are excluded from model access.')
paragraph('Publishing, restoring and unpublishing require a separate confirmation tied to an exact version or action, owner, project, expected publication revision and ten-minute expiry. D1 guards apply the pointer, activity event and chat receipt atomically. Draft work never changes the public pointer. Future purchases and destructive operations need equally exact impact and price approvals; ordinary language from a model is not authority.')

heading('28 AI Usage Credits and Payment Boundaries')
paragraph('The server adapter supports OpenAI Chat Completions and Responses text protocols. Operator configuration fixes the HTTPS endpoint, model, key and protocol. Prompts cannot change the destination. Requests use manual redirects, a 60-second timeout, bounded input/output and sanitized errors. Redirect responses are rejected. Automatic provider retries and fallback models are disabled.')
paragraph('The live test uses https://opencode.ai/zen/go/v1 and muse-spark-1.3-contributor through Responses, under the existing OpenCode Go subscription. Contributor data-use disclosure is shown to testers. Photo names and dimensions can be supplied to the planner; photo pixels are not sent, and the model must not invent visual descriptions. The provider boundary adapts reviewed Jarvis response validation and diagnostics; docs/jarvis-reuse.md records provenance and differences.')
table(['Plan', 'Monthly USD', 'Live sites', 'Monthly units'], [
 ('Explore', '0', '1', '25'), ('Launch', '12 proposed', '3', '250'), ('Studio', '29 proposed', '10', '800'),
])
paragraph('One build unit is the ceiling of (input tokens + four times output tokens) divided by 1,000. Provider reasoning tokens count as output usage. Reserve the bounded maximum before inference, save the answer and token receipt before settlement, then charge actual units and release the remainder. Monthly allowance is consumed before credits; credit overage requires explicit opt-in. Credits cannot bypass site, file, storage or abuse limits.')
paragraph('AI reservation and chat-turn claim commit atomically before dispatch. Saved token receipts reconcile after a lost acknowledgement without another inference. Confirmed failures release reservations; uncertain provider outcomes retain a hold. Undispatched reservations expire after five minutes. An operator-only reviewed waiver can absorb an unresolved operation cost after ten minutes, release its exact held quota/credits once, close the chat turn and preserve an immutable audit. Provider tokens remain unknown; the waiver does not assert zero provider usage. Late results cannot create drafts or settle a closed turn. Automated provider-receipt reconciliation remains outstanding. Guided commands and trusted component assembly consume no AI units.')
paragraph('The proposed top-up is USD 5 for 100 credits. Monthly units use UTC calendar months; credits persist separately. An internal plan-fulfillment boundary now requires immutable month-specific payment references. Upgrades raise the allowance ceiling without resetting spending or holds. New paid months need new receipts; otherwise the account uses Explore. Existing live sites remain available after a lower allowance, but new publication is blocked over the current tier limit. This internal function is not a public or model-grant endpoint.')

paragraph('Stripe test credit checkout has server-fixed orders, signed timestamped webhooks, authoritative session retrieval and idempotent grants. It remains unconfigured because the owner has no Stripe account. Subscriptions, refunds, disputes, proration, real checkout acceptance and real charges are unfinished. Domain registration is a separate future charge and cannot silently consume build credits.')

heading('29 Immutable Websites Photos and Exports')
paragraph('The generic catalog currently contains launch-page@1, site-page@1 and photo-page@1. Configuration is validated plain text, allowlisted themes, up to six text sections and at most four same-project ready photo references. The assembler escapes text and records exact HTML, configuration, component ID, digest, source and optional base version. D1 limits a project to 25 retained versions, 32 KiB HTML and 8 KiB configuration per version.')
paragraph('The first agent website action saves the current specification, rendered page and receipt atomically. Subsequent edits and photo changes require an exact private specification review and create a new immutable version. Owned approved photo dependencies commit with the reviewed version; same-owner cross-project photos are rejected. Publication remains a separate confirmation. No legacy builder, importer or slash-command interface is supported.')
sub('Private preview and public serving')
paragraph('Private previews require the current owner session and use no-store caching and a CSP sandbox that denies scripts and forms. Photo previews read and verify exact dependencies, then embed bounded JPEG data URLs into the response. Stored HTML retains relative asset paths. This avoids public bearer links and opaque-origin cookie dependencies.')
paragraph('Confirmed publication points to an exact retained version at PROJECT_ID.whumpo.site. Photo pages canonicalize a trailing slash. The public Worker checks the current publication relation before serving /s/PROJECT_ID/assets/FILE_ID.jpg. Unattached uploads and photos referenced only by other versions are inaccessible. Restore switches HTML and references together; unpublish revokes both on the next request. Publication reviews explicitly disclose photo exposure.')
sub('Uploads and retention')
paragraph('The browser accepts JPEG, PNG and WebP up to 10 MiB, decodes orientation, resizes to 1600 by 1600 and encodes a JPEG up to 1 MiB. The server validates the bounded baseline-JPEG envelope and strips APP/COM metadata; it is not a complete pixel decoder. D1 reserves immutable file identity before create-only R2 writes. Retries verify checksum and size; conflicting reuse is rejected. Pending reservations count toward limits and require operator recovery if the prepared upload bytes were lost.')
paragraph('Limits are 20 photos per project, 50 per account and 500 platform-wide, with queued physical deletions counted toward capacity. Each photo is at most 1 MiB. Atomic monthly counters cap application R2 reads at 100,000 and writes at 10,000; daily cleanup runs outside those counters. The included allowance is shared with other applications, so these controls are not an account-wide billing cap.')
paragraph('Source export returns exact HTML/CSS, referenced JPEGs, security headers, README and configuration/checksum metadata in a bounded ZIP. Missing or mismatched assets fail the export. Accounts, chats, keys, billing and database contents are excluded. Reviewed removal of an unused photo immediately removes access and queues R2 cleanup. Any retained version reference blocks removal, including a reference added after review. Applied-review tombstones prevent delayed uploads from recreating removed files.')

paragraph('Saved versions can be removed only after an exact ten-minute review and separate user confirmation. The transaction checks the owner, digest and expiry, deletes the version, invalidates its preview/export/restore links and records an audit event. Live versions, active publication reviews and all GitHub backup records retain their versions. Original photos and independent newer artifacts remain; photo references are released. Tombstones prevent old version IDs from being recreated. Removing backed-up versions, automatic pruning and whole-project/account deletion remain unfinished.')

heading('30 Shared Components and Backend Reuse')
paragraph('The first iteration of a customer site should reuse reviewed generic code whenever possible. Customer prompts, generated sites, uploads, business content, credentials and records must remain private. The shared catalog contains platform-authored source, appropriately licensed upstream code and synthetic fixtures. Customer customization is not a catalog contribution.')
paragraph('Each package version needs a stable ID, digest, provenance, license, configuration schema, runtime compatibility, dependencies, migrations and declared capabilities. Only catalog maintainers publish versions. Agents select exact approved IDs and validated parameters; download URLs, shell commands and executable configuration are not accepted. Immutable versions keep existing customer sites reproducible when the library changes.')
sub('Authentication and forms as the next backend components')
paragraph('A reusable sign-in component must separate constrained presentation from reviewed credential verification, sessions, recovery and access checks. Styling requests cannot alter authentication policy. Each instance resolves tenant identity from verified server context, not model-supplied owner IDs. Sharing implementation must never share accounts, sessions, secrets or customer database rows. Platform account authentication is implemented; authentication for customer-generated applications is not yet packaged.')
paragraph('The contact-request backend is deployed. Explicit activation pins the form to the current publication revision; publishing or restoring disables it until a fresh review. Bounded name/email/message input uses Origin, Turnstile and atomic rate/capacity checks. Messages are private to the owner, visible for 30 days and excluded from model context. The test caps intake at 20 per project per hour, 100 retained per project and 1,000 across the platform. The workspace inbox supports deletion; receipts prevent replay resurrection and quota recycling. No outbound email is sent. Successful human Turnstile completion and the first scheduled production retention run remain unverified.')
sub('Promotion review and caches')
paragraph('A customer adaptation stays private by default. Explicit consent permits nomination only. A separate synthetic candidate must remove customer configuration and identifiers, then pass human source review, secret scanning, dependency/license review and security/behavior checks. Pattern matching or a model assertion cannot prove that customer data is absent. Opaque assets and copied real-data fixtures are excluded.')
paragraph('Catalog retrieval searches only shared metadata and synthetic examples. Cache package bytes by immutable digest. Never cache customer prompts, filled configurations or completed sites globally. Private caches include owner and project boundaries and deletion semantics. Withdrawn components block new adoption and need a reviewed migration path for affected sites; do not silently rewrite customer deployments.')
paragraph('Measure component matches, avoided model calls, tokens per initial preview, later customization and defects. Guided static assembly and photo attachment already avoid inference. No percentage cost-saving claim is justified yet. The broader backend library, promotion workflow, component withdrawal and runtime isolation remain unfinished.')

sub('Directory data and controlled recovery')
paragraph('A directory holds up to 25 private listings per project, each with a title, description and optional credential-free HTTPS link. Add, edit and remove operations prepare exact reviews; only an explicit user confirmation changes the draft. A separate publication review exposes the complete snapshot through a responsive page and a read-only public JSON API. Draft edits and restoration leave the public snapshot unchanged. Parent website publication changes disable directory serving until another review. Shared code contains no customer records.')
paragraph('Applied edit/publication history records exact listing bytes and SHA-256. History reads, private JSON downloads and reviewed draft restoration are available through chat. Restore compares current and replacement records, retains listing IDs and records immutable source ancestry. Pending proposals, foreign projects and altered snapshots cannot be recovery sources. Source exports contain data and provenance, not a portable backend. The test retains 100 edit/publication reviews per project; reaching that limit does not prevent taking the directory offline. Automatic history pruning and generic application migrations remain unfinished.')
sub('Version comparisons and publication health')
paragraph('Private version comparisons pin owned same-project targets and baselines, showing supported configuration, section and photo-reference changes. They are not pixel or arbitrary source-code diffs. Publication reviews can expose their exact comparison; stale reviews are refused.')
paragraph('Publishing and restoring attempt an HTTP check of the shared published page. The checker verifies stored artifact integrity, exact served HTML including the current contact link, and up to four photo HEAD responses. It follows no redirects, accepts no arbitrary URL, and sends no account credentials. Results are fenced by publication/contact revisions and carry timestamps. Custom domains, visitor submissions, browser behavior and image decoding are outside this check.')
paragraph('Database claims limit a given publication/contact state to one check per five minutes and all checks to 500 per UTC day. The hourly job considers 20 due sites in batches of four; healthy results older than six hours are labeled stale. Check history is retained for 30 days. A failed check records an unhealthy result without automatic rollback; a real scheduled production run remains unobserved.')

heading('31 Domains Connections and Future Workloads')
paragraph('Existing-domain ownership is implemented through owner-scoped hostname claims, random TXT proofs, seven-day expiry and bounded fixed-resolver checks. The agent and guided commands can prepare, renew and check proofs. Exact positive matching verifies ownership only. It does not connect hosting, issue a certificate or imply domain registration. Project/account/platform limits are three, ten and eighty claims respectively.')
paragraph('Custom-hostname hosting code reviews and provisions exact provider resources, verifies current ownership and direct CNAME routing, and requires active hostname/TLS state before granting a short serving lease. Refresh results are fenced by attempt ID and current proof. Unknown creation outcomes never trigger a blind second POST or adoption by name. Disconnect removes the serving lease before exact-ID cleanup. Operator recovery can reconcile a known provider ID only with independent evidence and an immutable audit receipt. This code remains disabled pending scoped platform credentials and verified fallback-origin setup; no real customer hostname or certificate has been provisioned.')
paragraph('Domain purchase is a separate limited capability. Search and availability are advisory until an authoritative price quote identifies registration term, total charge, renewal terms, supported extension and expiry. Purchase requires exact explicit confirmation and an idempotent operation record. Unknown registrar results must be reconciled before another attempt. Registrar access, payments and customer terms are not configured; do not simulate successful registration or charge build credits.')
paragraph('Optional GitHub App and Cloudflare OAuth code uses one-time session-bound S256 PKCE, encrypted expiring credentials, scoped authorization, immediate local disconnect and leased remote revocation. AES-GCM binds provider, owner, record and purpose under a separate vault key. No operator GitHub or Wrangler token is used as a customer credential. The flows remain disabled because client/App credentials and the vault key are absent. Private Cloudflare clients initially accept only parent-account members; external testers require a separate public-client review.')
sub('Reviewed source backups and customer DNS')
paragraph('GitHub backup reviews pin a saved version and source hash, selected repository visibility/branch and expected branch head. A confirmed commit writes a readable version folder with an expected-head condition and skip-ci message. Exact tree checks verify the result. Uncertain writes are reconciled read-only, never resent. The implementation requires an existing initialized repository; real consent and backup commits remain unverified without App setup.')
paragraph('Customer DNS reviews pin an owned claim, active authorized zone, exact TXT proof, credential ID, complete record inventory and ten-minute expiry. Identical existing proofs need no write. Only explicit user confirmation can add the TXT record; current records are never overwritten. A committed dispatch state and unique attempt guards prevent duplicate writes. An uncertain outcome is only observed, and a matching marker does not grant deletion authority. Public DNS verification remains separate.')
paragraph('Chat now includes DNS zone discovery, TXT proof reviews, fixed-target routing CNAME reviews, hosting review/confirm/check/disconnect, and GitHub installation/repository discovery and source-backup review/confirm/check. Models can prepare the permitted reviews, never confirm them. The server claims a durable chat turn before provider work and supplies the actual receipt. Fixed-target routing refuses conflicting records and requires explicit confirmation; uncertain results are only observed. UI setup remains disabled without the corresponding credentials. Exact-record editing/removal, long-lived renewal and registrar operations are unfinished.')

paragraph('The directory starter now supplies a bounded dynamic CRUD service using trusted shared Worker code and project-scoped D1 data. General schemas, scheduled customer jobs and isolated generated build execution remain open. The free/included-service constraint prevents adopting Containers or Sandbox solely because the original proposal named them. A future arbitrary-code runtime needs a verified isolation boundary, resource limits and a compatible cost model before customer use.')
paragraph('Traditional VPS workloads, fixed-machine semantics and email mailbox hosting remain outside the intended product. Website version comparisons, timestamped shared-address health checks and directory data restoration are now delivered for their supported schemas. Durable multi-step job orchestration, automatic rollback and general database/service recovery remain outstanding. Existing safe publication and restoration do not prove those broader capabilities.')

heading('32 Verification Operations and Remaining Gates')
paragraph('Deploy with tools/deploy.mjs: apply remote D1 migrations, build assets, deploy the public host, then deploy the workspace Worker. A migration failure stops deployment. An earlier trigger syntax accepted locally failed remotely; the deployment gate and compatible trigger form were corrected. Record both Worker versions and migration state in docs/deployment.md. Do not run paid CI or enable additional provider services automatically.')
table(['Area', 'Evidence and limitation'], [
 ('Identity and chat', 'Local isolation/concurrency/replay tests and live signed passkey assertions pass. Native synced-device acceptance remains.'),
 ('AI and metering', 'Live Muse creation/editing, inspection and DNS-review discovery pass. Atomic month receipts and concurrent claims pass controlled tests. Unknown provider usage still needs reconciliation; audited cost waivers are available.'),
 ('Versions and photos', 'Local ownership, cross-project rejection, restore/unpublish, archive integrity and quota concurrency pass. Live publication and browser photo rendering pass.'),
 ('Export', 'Private directory JSON preserves exact saved listings and digest. Independent ZIP CRC and image decoding pass. Browser policy blocked opening local file URLs; direct offline browser rendering is unverified.'),
 ('Retention', 'Photo and version removal tests cover review races, replay, expiry and lost acknowledgements. Live individual-removal checks pass. Whole-account lifecycle and first scheduled production cleanup observation remain open.'),
 ('Payments and integrations', 'Controlled OAuth, backup, TXT/CNAME and custom-hostname tests pass. Live disabled-access checks pass. Real customer consent, provider writes, TLS and checkout acceptance remain pending configuration.'),
])
paragraph('Application snapshot: source commit 771facf; migrations through 0031. Workspace Worker 816e7512-a60e-48a5-87dc-50e2b48c9a3f and public Worker c7b5adb3-18e2-4961-a7e4-9211bbf051c3. Daily 03:17 UTC maintenance handles contacts, photos and provider revocation; hourly minute 47 handles hostname maintenance and bounded site-health refresh. This document revision does not imply a runtime deployment. A read-only live secret inventory confirmed only AI_API_KEY, PASSWORD_PEPPER and TURNSTILE_SECRET are installed; customer GitHub, Cloudflare OAuth and custom-hostname setup are still absent.')
paragraph('Operations still need transcript/audit retention, customer project/account deletion, abuse suspension, pending-upload recovery, provider-usage reconciliation and tested disaster recovery. Private directory data export is now available alongside static source export. Individual unused photo/version removal and directory draft deletion do not constitute whole-account erasure. Provider revocation and external-resource audit records must be accounted for before whole-account deletion. R2 cleanup remains eventual, and account-wide usage is separate from application caps.')
paragraph('The next delivery gates are a fresh-account natural-language creation plus three edits within the published allowance; real customer hostname/TLS and optional-account acceptance after setup; registrar quotes and exact purchase confirmation if that beta is enabled; broader dynamic workloads and isolated execution within the cost constraint; and device, Turnstile, mobile and operations acceptance. Explore currently provides 25 units; separate successful model calls or tests with injected credits do not prove the complete three-edit journey fits that allowance. Payment setup remains deferred by the owner. These milestones do not complete the full MVP.')
paragraph('Detailed current contracts are in docs/mvp.md, docs/version-retention.md, docs/photos.md, docs/plan-lifecycle.md, docs/payments.md, docs/shared-components.md, docs/workspace-agent.md, docs/deployment.md and docs/integrations. Original product intent remains in docs/reference/ai_managed_cloud_mvp_design.docx. Appendix product notes retain their historical date and are not current eligibility or pricing guarantees. Recheck provider capabilities before enabling an integration.')

heading('33 Acceptance Evidence and Remaining Work')
paragraph('This audit preserves the original acceptance criteria and the owner’s later constraints. Delivered means the supported path has source and recorded verification evidence. It does not imply every device, provider account, workload or failure mode was tested. Unverified and setup-gated items remain incomplete. The detailed evidence map is docs/acceptance.md.')
table(['Original acceptance area', 'Current determination'], [
 ('Registration and project start', 'Implemented without external customer accounts. Signed assertion and guided onboarding checks pass. Real device passkey and successful human signup challenge remain unverified.'),
 ('Natural-language preview and three edits', 'Muse preview/edit paths work. A fresh Explore account completing creation plus three subsequent edits under its normal 25-unit quota still needs one end-to-end run.'),
 ('Publish with health and restore', 'Exact version publication, shared-address HTTP receipts, private comparison and confirmed restore are deployed. Automatic rollback and custom-domain health remain open.'),
 ('Uploads and incorporation', 'Bounded private JPEG photos and exact same-project attachment are delivered. Arbitrary file types, generated assets and general build artifacts are not supported.'),
 ('Existing domain with HTTPS', 'Proof, routing review and custom-hostname lifecycle code exist. Live SaaS credentials/origin setup and real TLS/routing acceptance are missing.'),
 ('Domain purchase when enabled', 'Disabled and unimplemented. Requires supported registrar access, authoritative terms/price, funding and explicit confirmation. No purchase or fake receipt is enabled.'),
 ('Failed build isolation', 'Rejected template configuration cannot replace an immutable published artifact. General build/test isolation and known-good automated recovery are not proven.'),
 ('Identity audit and credentials', 'Owner joins, Origin checks, guided confirmation and credential boundaries are implemented. Cross-tenant tests cover supported resources. Whole-product privileged-action traceability still requires a consolidated audit.'),
 ('Default customer experience', 'Chat is primary; details panels and provider consent are secondary. Free guided alternatives exist. Dense review IDs/digests and full mobile/device journeys need further UX acceptance.'),
 ('Plans and credits', 'Proposed tiers, atomic reservations, month receipts, opt-in credits and audited waivers exist. Checkout/paid lifecycle are deferred or unconfigured; no real charges are enabled.'),
 ('Dynamic services and operations', 'Contact inbox and directory CRUD/recovery are live. Generic schemas, isolated execution, jobs, account deletion, abuse operations and disaster recovery remain open.'),
])
paragraph('Provider setup is a concrete external dependency: the workspace has no GitHub App config, Cloudflare OAuth client config or custom-hostname token file, and the live Worker secret inventory lacks their credentials. Native passkey and human challenge acceptance require user-device interaction. Continue independent implementation while those gates remain pending; do not count disabled-path tests as real-provider acceptance.')

import re
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        p.text = re.sub(r'[^A-Za-z0-9 ]+', ' ', p.text)
        p.text = re.sub(r' +', ' ', p.text).strip()

from docx.oxml import OxmlElement
for t in doc.tables:
    if len(t.rows) > 1:
        repeat = OxmlElement('w:tblHeader')
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for c in t.rows[0].cells:
            for p in c.paragraphs: p.paragraph_format.keep_with_next = True

for section in doc.sections:
    for p in section.footer.paragraphs:
        if p.text.strip(): p.text = 'AI Managed Cloud Platform MVP Design  |  v0.5'
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    doc.styles[name].font.color.rgb = RGBColor(0,0,0)
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
